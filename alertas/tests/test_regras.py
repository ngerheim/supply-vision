import numpy as np
import pandas as pd
from openpyxl import load_workbook


KEY = ["_fornec_norm", "_cidade_norm", "_modelo_norm", "_peca_norm"]


def acordo(precos):
    linhas = []
    for preco in precos:
        linhas.append({
            "_fornec_norm": "1", "_cidade_norm": "X", "_modelo_norm": "M",
            "_peca_norm": "ITEM", "PECA_SERVICO": "ITEM", "FORNECEDOR": "F",
            "PRECO": preco, "_preco_original": str(preco),
            "_preco_valido": bool(pd.notna(preco) and np.isfinite(preco) and preco >= 0),
        })
    return pd.DataFrame(linhas)


def base(qtd=1, preco=100):
    return pd.DataFrame({
        "Modelo": ["M"] * qtd, "_desc_norm": ["ITEM"] * qtd,
        "_cidade_norm": ["X"] * qtd, "_fornec_norm": ["1"] * qtd,
        "Valor Unitario": [preco] * qtd, "OS Quantidade": [1] * qtd,
        "Data Abertura": ["01/08/2026"] * qtd, "Codigo OS": range(1, qtd + 1),
        "Forncedor por Cidade": ["X"] * qtd, "Fornecedor": ["F"] * qtd,
        "Descrição": ["ITEM"] * qtd, "Criado Por": ["T"] * qtd,
        "Grupo Despesa": ["G"] * qtd,
    })


def com_vigencia(df, inicio="18/09/2026", fim=None, status="active"):
    df = df.copy()
    df["INICIO_VIGENCIA"] = inicio
    df["FIM_VIGENCIA"] = fim
    df["STATUS_ACORDO"] = status
    return df


def test_antes_do_corte_todos_os_acordos_sao_considerados_ativos(rodar):
    compras = base(preco=10)
    compras["Data Abertura"] = "17/09/2026"
    futuro_suspenso = com_vigencia(acordo([10]), inicio="01/12/2026", status="suspended")
    assert rodar.processar(compras, futuro_suspenso).loc[0, "Status"] == "CONFORME"


def test_a_partir_do_corte_respeita_status_e_inicio_da_vigencia(rodar):
    compras = base(qtd=3, preco=10)
    compras["Data Abertura"] = ["18/09/2026", "30/09/2026", "01/10/2026"]
    vigente = com_vigencia(acordo([10]), inicio="30/09/2026", status="active")
    resultado = rodar.processar(compras, vigente)
    assert resultado["Status"].tolist() == ["SEM ACORDO", "CONFORME", "CONFORME"]


def test_fim_da_vigencia_e_inclusivo(rodar):
    compras = base(qtd=2, preco=10)
    compras["Data Abertura"] = ["30/09/2026", "01/10/2026"]
    vigente = com_vigencia(acordo([10]), inicio="18/09/2026", fim="30/09/2026")
    resultado = rodar.processar(compras, vigente)
    assert resultado["Status"].tolist() == ["CONFORME", "SEM ACORDO"]


def test_precos_iguais_na_precisao_monetaria(rodar):
    ac = acordo([10.0, 10.0000000001])
    ac["PRECO"] = pd.to_numeric(ac["PRECO"]).round(2)
    resultado = rodar.processar(base(preco=10), ac)
    assert resultado.loc[0, "Status"] == "CONFORME"


def test_carregamento_arredonda_e_rejeita_preco_nao_finito(rodar):
    bruto = pd.DataFrame({
        "MODELO": ["M"] * 5, "PECA_SERVICO": ["ITEM"] * 5,
        "CIDADE": ["X"] * 5, "CNPJ": ["1"] * 5,
        "PRECO": [10, 10.0000000001, 0, -1, np.inf], "FORNECEDOR": ["F"] * 5,
    })
    carregado = rodar._preparar_acordos(bruto)
    assert carregado.loc[0, "PRECO"] == carregado.loc[1, "PRECO"] == 10.0
    assert carregado["_preco_valido"].tolist() == [True, True, True, False, False]


def test_celula_vazia_de_preco_nao_vira_cortesia(rodar):
    bruto = pd.DataFrame({
        "MODELO": ["M"], "PECA_SERVICO": ["ITEM"], "CIDADE": ["X"],
        "CNPJ": ["1"], "PRECO": [""], "FORNECEDOR": ["F"],
    })
    carregado = rodar._preparar_acordos(bruto)
    assert not bool(carregado.loc[0, "_preco_valido"])


def test_ambiguo_nao_expoe_nenhuma_referencia(rodar):
    resultado = rodar.processar(base(), acordo([10.0, 20.0]))
    assert resultado.loc[0, "Status"] == rodar.STATUS_AMBIGUO
    for coluna in ("Preco Acordo", "Preco Total Acordo", "Diferenca Unit.",
                   "Diferenca Total", "Menor Preco Acordo", "Dif. p/ Menor Acordo"):
        assert pd.isna(resultado.loc[0, coluna])
    assert resultado.loc[0, "Fornecedor do Acordo"] == ""
    assert resultado.loc[0, "Tinha acordo?"] == ""


def test_preco_nulo_negativo_ou_infinito_gera_pendencia(rodar):
    for preco in (np.nan, -1, np.inf):
        resultado = rodar.processar(base(), acordo([preco]))
        assert resultado.loc[0, "Status"] == rodar.STATUS_PRECO_INVALIDO


def test_cortesia_cobrada_sai_como_acima_do_acordo(rodar):
    """Acordo com preço 0 é cortesia: compra cobrada é desvio, não pendência."""
    resultado = rodar.processar(base(preco=30), acordo([0.0]))
    assert resultado.loc[0, "Status"] == "ACIMA DO ACORDO"
    assert resultado.loc[0, "Preco Acordo"] == 0.0
    assert resultado.loc[0, "Diferenca Unit."] == 30.0
    assert resultado.loc[0, "Diferenca Total"] == 30.0
    assert resultado.loc[0, "Menor Preco Acordo"] == 0.0


def test_cortesia_ao_lado_de_preco_positivo_e_ambigua(rodar):
    """Mesma chave com 0 e preço positivo: qual vale é indeterminado."""
    resultado = rodar.processar(base(), acordo([0.0, 50.0]))
    assert resultado.loc[0, "Status"] == rodar.STATUS_AMBIGUO


def test_denominador_exclui_quarentena(rodar):
    df = pd.DataFrame({"Status": ["CONFORME"] * 10 + [rodar.STATUS_AMBIGUO] * 90,
                       "Motivo Sem Acordo": [""] * 100})
    resumo = rodar.resumir_status(df)
    assert resumo["total_bruto"] == 100
    assert resumo["total_elegivel"] == 10
    assert resumo["percentuais_elegiveis"]["CONFORME"] == 100.0
    assert sum(resumo["contagens"].values()) == 100


def test_tres_grupos_sao_mutuamente_exclusivos(rodar):
    statuses = pd.Series(["CONFORME", "ACIMA DO ACORDO", "ABAIXO DO ACORDO",
                          "SEM ACORDO", rodar.STATUS_AMBIGUO, rodar.STATUS_PRECO_INVALIDO])
    grupos = [statuses.isin(rodar.STATUS_COM_ACORDO), statuses == "SEM ACORDO",
              statuses.isin(rodar.STATUS_QUARENTENA)]
    assert sum(int(g.sum()) for g in grupos) == len(statuses)
    assert all(sum(bool(g.iloc[i]) for g in grupos) == 1 for i in range(len(statuses)))


def test_nomes_de_modelo_e_item_sao_traduzidos_com_fallback(rodar):
    rodar.MODELOS.update({"M": "MODELO PADRONIZADO"})
    rodar.SINONIMOS.update({"ITEM": "ITEM PADRONIZADO"})
    resultado = rodar.processar(base(qtd=2), acordo([90]))
    assert resultado.loc[0, "Modelo"] == "MODELO PADRONIZADO"
    assert resultado.loc[0, "Item"] == "ITEM PADRONIZADO"

    original = base()
    original["Modelo"] = "SEM DE-PARA"
    original["Descrição"] = "DESCRICAO ORIGINAL"
    original["_desc_norm"] = "DESCRICAO ORIGINAL"
    sem_mapa = rodar.processar(original, acordo([90]))
    assert sem_mapa.loc[0, "Modelo"] == "SEM DE-PARA"
    assert sem_mapa.loc[0, "Item"] == "DESCRICAO ORIGINAL"


def test_anexo_diario_reproduz_layout_classico_sem_verde(rodar, tmp_path):
    dados = rodar.processar(base(preco=100), acordo([90]))
    destino = tmp_path / "divergencias.xlsx"
    rodar.gerar_alerta_acordo(dados, destino)

    wb = load_workbook(destino)
    ws = wb["Relatório"]
    cabecalhos = [celula.value for celula in ws[1]]
    coluna_cnpj = cabecalhos.index("CNPJ") + 1
    letra_cnpj = ws.cell(1, coluna_cnpj).column_letter
    assert ws.column_dimensions[letra_cnpj].hidden is not True
    assert ws.cell(2, coluna_cnpj).value == "00000000000001"
    assert len(ws.conditional_formatting) == 4
    formulas = {
        formula
        for regras in ws.conditional_formatting._cf_rules.values()
        for regra in regras
        if regra.type == "expression"
        for formula in regra.formula
    }
    assert formulas == {
        '$O2="ACIMA DO ACORDO"',
        '$O2="ABAIXO DO ACORDO"',
    }
    cores_linhas = {
        regra.dxf.fill.bgColor.rgb
        for regras in ws.conditional_formatting._cf_rules.values()
        for regra in regras
        if regra.type == "expression"
    }
    assert cores_linhas == {"FFFDEAEA", "FFEAF1FB"}
    assert not any("CONFORME" in formula for formula in formulas)
    assert ws.cell(2, cabecalhos.index("Item") + 1).alignment.wrap_text is True
    assert ws.cell(1, 1).fill.fgColor.rgb == "FF52606E"
    assert ws.cell(1, cabecalhos.index("Item") + 1).fill.fgColor.rgb == "FF4C5578"
    assert ws.cell(1, cabecalhos.index("Preco OS") + 1).fill.fgColor.rgb == "FF3F6B6B"
    tabela = next(iter(ws.tables.values()))
    assert tabela.tableStyleInfo.showRowStripes is False
    assert tabela.tableStyleInfo.name == "TableStyleLight1"


def test_filtro_do_alerta_exclui_conformes(rodar):
    statuses = pd.Series(["CONFORME", "ACIMA DO ACORDO", "ABAIXO DO ACORDO", "SEM ACORDO"])
    assert statuses[statuses.isin(rodar.STATUS_DIVERGENCIA)].tolist() == [
        "ACIMA DO ACORDO", "ABAIXO DO ACORDO"
    ]


def test_recorte_historico_reune_quatro_estados_e_indica_acordo(rodar):
    entrada = pd.DataFrame({
        "Status": ["CONFORME", "ACIMA DO ACORDO", "ABAIXO DO ACORDO",
                   "SEM ACORDO", rodar.STATUS_AMBIGUO],
        "Item": ["A", "B", "C", "D", "E"],
    })
    dados = rodar.preparar_recorte_historico(entrada)
    assert dados["Status"].tolist() == [
        "CONFORME", "ACIMA DO ACORDO", "ABAIXO DO ACORDO", "SEM ACORDO"
    ]
    assert dados["Com acordo"].tolist() == ["SIM", "SIM", "SIM", "NÃO"]


def test_planilha_historica_e_unica_e_contem_coluna_com_acordo(rodar, tmp_path):
    partes = [
        rodar.processar(base(preco=90), acordo([90])),
        rodar.processar(base(preco=100), acordo([90])),
        rodar.processar(base(preco=80), acordo([90])),
        rodar.processar(base(preco=90), acordo([90]).assign(_fornec_norm="OUTRO")),
    ]
    destino = tmp_path / "recorte.xlsx"
    assert rodar.gerar_recorte_historico(pd.concat(partes, ignore_index=True), destino)
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Relatório"]
    ws = wb.active
    cabecalhos = [celula.value for celula in ws[1]]
    assert "Com acordo" in cabecalhos
    coluna = cabecalhos.index("Com acordo") + 1
    assert [ws.cell(r, coluna).value for r in range(2, 6)] == ["SIM", "SIM", "SIM", "NÃO"]
    assert len(ws.conditional_formatting) == 0
    assert all(celula.alignment.wrap_text is not True for linha in ws for celula in linha)
    tabela = next(iter(ws.tables.values()))
    assert tabela.tableStyleInfo.showRowStripes is True
    assert tabela.tableStyleInfo.name == "TableStyleMedium15"
    coluna_cnpj = cabecalhos.index("CNPJ") + 1
    assert ws.cell(2, coluna_cnpj).value == "00000000000001"
    assert ws.row_dimensions[1].height == 20
    assert ws.sheet_format.defaultRowHeight == 18
    coluna_item = ws.cell(1, cabecalhos.index("Item") + 1).column_letter
    coluna_fornecedor = ws.cell(1, cabecalhos.index("Fornecedor") + 1).column_letter
    assert ws.column_dimensions[coluna_item].width >= 42
    assert ws.column_dimensions[coluna_fornecedor].width >= 26
